import base64
import json
from io import BytesIO

from ajax_helpers.utils import ajax_command
from django.http import QueryDict
from django.utils.html import strip_tags
from django_menus.menu import MenuItem
from openpyxl import Workbook
from openpyxl.cell import Cell

from django_datatables.helpers import add_filters


class ExcelDownload:
    tables: dict

    ajax_commands = ['column']
    excel_filename = 'download.xlsx'
    excel_id = 'id'

    @staticmethod
    def sort_excel(table, table_array):
        # Hook: subclasses can reorder / sort the exported rows before they are written.
        return table_array

    def check_pk_column(self, table):
        # The exported id column defaults to 'id'; use the model's real pk column when it differs.
        if table.model is not None and table.model._meta.pk.column != 'id':
            self.excel_id = table.model._meta.pk.column

    def download_menu_item(self, table_name=None):
        if table_name is None:
            table_name = list(self.tables.keys())[0]
        return MenuItem(ajax_command('send_column', method='get_excel', column='id', table_id=table_name),
                        'Download', font_awesome='fas fa-file-excel', link_type=MenuItem.AJAX_COMMAND)

    # noinspection PyUnresolvedReferences
    def button_download_all(self):
        table = list(self.tables.values())[0]
        self.setup_tables(table_id=table.table_id)
        self.excel_table(table)
        self.check_pk_column(table)
        return self.download_excel(table)

    def download_excel(self, table, query=None):
        workbook = Workbook()
        sheet = workbook.active

        col_format = []
        excel_styles = []
        titles = []

        for n, c in enumerate(table.columns):
            if getattr(c, 'xl_dont_show', lambda: False)() or c.options.get('hidden'):
                col_format.append(False)
            else:
                titles.append(str(strip_tags(c.title)))
                col_format.append(c.excel if hasattr(c, 'excel') else True)
                if hasattr(c, 'xl_style'):
                    excel_styles.append((n, c.xl_style))

        sheet.append(titles)
        if query is None:
            query = table.table_data if table.table_data else self.get_table_query(table)
        results = self.sort_excel(table, table.get_table_array(self.request, query))
        for r in results:
            row = [Cell(value=(d(v) if d != True else v), worksheet=sheet) for v, d in zip(r, col_format) if d]
            for f in excel_styles:
                f[1](row[f[0]])
            sheet.append(row)
        for col in sheet.columns:
            column = col[0].column_letter
            sheet.column_dimensions[column].width = 10
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        # noinspection PyUnresolvedReferences
        return self.command_response('save_file', data=base64.b64encode(output.read()).decode('ascii'),
                                     filename=self.excel_filename)

    def excel_table(self, table):
        # can modify table when exporting to Excel. e.g. hide columns
        pass

    # noinspection PyUnresolvedReferences
    def column_get_excel(self, **kwargs):
        table = self.tables[kwargs['table_id']]
        self.setup_tables(table_id=table.table_id)
        self.excel_table(table)
        self.check_pk_column(table)
        from django_datatables.datatables.server_side import ServerSideTable
        if isinstance(table, ServerSideTable) and kwargs.get('datatable_state') is not None:
            # The browser only holds the current page, so instead of a list of
            # ids the client sends its last DataTables request state; rebuild
            # the full filtered queryset from it.
            state = QueryDict(kwargs['datatable_state'])
            query = table.filtered_query(state, self.get_table_query(table))
            return self.download_excel(table, query=query)
        column_data = json.loads(kwargs['column_data'])
        table.filter = add_filters(table.filter, {f'{self.excel_id}__in': column_data})
        return self.download_excel(table)
